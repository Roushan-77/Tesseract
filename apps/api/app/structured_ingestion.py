import csv
import json
import re
import xml.etree.ElementTree as ET
from datetime import date, datetime
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

from .prompt4 import extract_relations_and_events

FIELD_ALIASES = {
    "PERSON": {"person", "name", "full_name", "username", "user", "driver", "operator", "complainant", "accused", "owner", "target_subject"},
    "PHONE": {"phone", "phone_number", "mobile", "caller", "receiver", "from_phone", "to_phone", "calling_number", "called_number", "sender_phone", "recipient_phone"},
    "EMAIL": {"email", "email_address"},
    "LOCATION": {"location", "address", "cell_tower", "cell tower", "place", "checkpoint", "toll_plaza", "toll_booth", "warehouse", "destination", "origin", "site", "target_location"},
    "VEHICLE": {"vehicle", "vehicle_number", "vehicle_no", "registration", "vehicle_reg", "truck_number", "truck_no", "plate_number", "license_plate"},
    "ORGANIZATION": {"organization", "organisation", "company", "employer", "bank", "vendor", "agency", "firm", "intermediary_bank"},
    "CASE_ID": {"case_id", "case", "fir", "fir_no", "record_id", "observation_id"},
    "DATETIME": {"date", "time", "timestamp", "datetime", "event_time", "transaction_date", "call_date", "movement_time", "observed_at", "pass_time"},
    "ACCOUNT": {"account", "account_id", "sender_account", "receiver_account", "from_account", "to_account", "acc_no", "account_number"},
    "AMOUNT": {"amount", "value", "transaction_amount", "sum", "total", "inr"},
}


def _clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return re.sub(r"\s+", " ", str(value)).strip()


def _field_type(name: str) -> str | None:
    normalized = re.sub(r"[^a-z0-9]+", "_", name.casefold()).strip("_")
    for field_type, aliases in FIELD_ALIASES.items():
        if normalized in aliases:
            return field_type
    if "duration" in normalized:
        return "CALL_DURATION"
    if "transaction" in normalized and "id" in normalized:
        return "TRANSACTION_ID"
    if "message" in normalized or "post" in normalized or "description" in normalized or "observation" in normalized or "activity" in normalized:
        return "TEXT"
    return None


def _records_from_csv(path: Path) -> tuple[list[dict[str, Any]], list[str], list[str]]:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as source:
            sample = source.read(4096)
            source.seek(0)
            dialect = csv.Sniffer().sniff(sample) if sample.strip() else csv.excel
            reader = csv.DictReader(source, dialect=dialect)
            columns = list(reader.fieldnames or [])
            return list(reader), [], columns
    except (OSError, UnicodeDecodeError, csv.Error) as exc:
        raise ValueError(f"CSV parsing failed: {exc}") from exc


def _records_from_workbook(path: Path) -> tuple[list[dict[str, Any]], list[str], list[str], dict[str, Any]]:
    warnings: list[str] = []
    sheets_data: dict[str, Any] = {}
    if path.suffix.lower() == ".xls":
        import xlrd
        book = xlrd.open_workbook(path)
        records = []
        all_columns = []
        for sheet in book.sheets():
            headers = [_clean(value) for value in sheet.row_values(0)]
            all_columns.extend(headers)
            sheet_rows = []
            for row_number in range(1, sheet.nrows):
                row_vals = sheet.row_values(row_number)
                sheet_rows.append(row_vals)
                records.append({f"{sheet.name}.{header}": sheet.cell_value(row_number, index) for index, header in enumerate(headers) if header})
            sheets_data[sheet.name] = {"columns": headers, "rows": sheet_rows}
        return records, warnings, list(dict.fromkeys(all_columns)), sheets_data
    workbook = load_workbook(path, read_only=True, data_only=True)
    records = []
    all_columns = []
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            headers = [_clean(value) for value in next(rows, ())]
            all_columns.extend(headers)
            sheet_rows = []
            for row in rows:
                sheet_rows.append(list(row))
                records.append({f"{sheet.title}.{header}": value for header, value in zip(headers, row) if header})
            sheets_data[sheet.title] = {"columns": headers, "rows": sheet_rows}
    finally:
        workbook.close()
    return records, warnings, list(dict.fromkeys(all_columns)), sheets_data


def _records_from_json(path: Path) -> tuple[list[dict[str, Any]], list[str], list[str]]:
    try:
        value = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValueError(f"JSON parsing failed: {exc}") from exc
    if isinstance(value, list) and all(isinstance(item, dict) for item in value):
        columns = list(dict.fromkeys(k for item in value for k in item.keys()))
        return value, [], columns
    if isinstance(value, dict):
        for key in ("records", "data", "rows", "items", "results"):
            candidate = value.get(key)
            if isinstance(candidate, list) and all(isinstance(item, dict) for item in candidate):
                columns = list(dict.fromkeys(k for item in candidate for k in item.keys()))
                return candidate, [], columns
    raise ValueError("JSON must contain an array of record objects")


def _records_from_xml(path: Path) -> tuple[list[dict[str, Any]], list[str], list[str]]:
    try:
        root = ET.parse(path).getroot()
    except (OSError, ET.ParseError) as exc:
        raise ValueError(f"XML parsing failed: {exc}") from exc
    records = []
    children = list(root)
    candidates = children if children and all(list(child) for child in children) else [root]
    for record in candidates:
        values = {_clean(child.tag): _clean(child.text) for child in list(record)}
        if values:
            records.append(values)
    if not records:
        raise ValueError("XML did not contain repeated record elements")
    columns = list(dict.fromkeys(k for item in records for k in item.keys()))
    return records, [], columns


def _mentions(records: list[dict[str, Any]], columns: list[str], sheets_data: dict[str, Any] | None = None) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    entities: list[dict[str, Any]] = []
    relations: list[dict[str, Any]] = []
    events: list[dict[str, Any]] = []
    normalized_records = []
    mention_id = 1
    
    for row_number, record in enumerate(records, start=2):
        typed: dict[str, list[dict[str, Any]]] = {}
        normalized = {str(key): _clean(value) for key, value in record.items() if _clean(value)}
        normalized_records.append({"row": row_number, "fields": normalized})
        
        for field, value in normalized.items():
            clean_field_name = field.rsplit(".", 1)[-1]
            field_type = _field_type(clean_field_name)
            if not field_type or field_type in {"DATETIME", "AMOUNT", "CALL_DURATION", "TRANSACTION_ID", "TEXT"}:
                continue
            entity_type = "PERSON" if field_type == "PERSON" else field_type
            mention = {
                "mentionId": f"structured-{row_number}-{mention_id}",
                "type": entity_type,
                "text": value,
                "normalizedValue": value,
                "start": row_number,
                "end": row_number,
                "confidence": 0.98,
                "sourceRow": row_number,
                "sourceField": field
            }
            mention_id += 1
            entities.append(mention)
            typed.setdefault(entity_type, []).append(mention)

        source_text = json.dumps(normalized, ensure_ascii=False)
        timestamp = next((value for field, value in normalized.items() if _field_type(field.rsplit(".", 1)[-1]) == "DATETIME"), None)

        # 1. CDR: Caller & Receiver Phones
        caller = next((value for field, value in normalized.items() if _field_type(field.rsplit(".", 1)[-1]) == "PHONE" and ("caller" in field.casefold() or "from" in field.casefold())), None)
        receiver = next((value for field, value in normalized.items() if _field_type(field.rsplit(".", 1)[-1]) == "PHONE" and ("receiver" in field.casefold() or "to" in field.casefold())), None)
        tower_loc = typed.get("LOCATION", [None])[0] if typed.get("LOCATION") else None
        
        if caller and receiver and typed.get("PHONE"):
            src_candidates = [item for item in typed["PHONE"] if item["text"] == caller]
            tgt_candidates = [item for item in typed["PHONE"] if item["text"] == receiver]
            if src_candidates and tgt_candidates:
                source, target = src_candidates[0], tgt_candidates[0]
                relations.append({
                    "type": "CONTACTED",
                    "source": source,
                    "target": target,
                    "sourceText": source_text,
                    "confidence": 1.0,
                    "sourceRow": row_number,
                    "sourceFields": [f for f in normalized if "caller" in f.casefold() or "receiver" in f.casefold() or "from" in f.casefold() or "to" in f.casefold()]
                })
                events.append({
                    "type": "CALL",
                    "date": timestamp,
                    "location": tower_loc,
                    "participants": [source, target],
                    "sourceText": source_text,
                    "confidence": 1.0,
                    "sourceRow": row_number
                })

        # 2. Financial: Sender & Receiver Accounts
        sender = next((value for field, value in normalized.items() if "sender_account" in field.casefold() or "from_account" in field.casefold() or "sender" in field.casefold()), None)
        receiver_account = next((value for field, value in normalized.items() if "receiver_account" in field.casefold() or "to_account" in field.casefold() or "receiver" in field.casefold() or "beneficiary" in field.casefold()), None)
        
        if sender and receiver_account and typed.get("ACCOUNT"):
            src_candidates = [item for item in typed["ACCOUNT"] if item["text"] == sender]
            tgt_candidates = [item for item in typed["ACCOUNT"] if item["text"] == receiver_account]
            if src_candidates and tgt_candidates:
                source, target = src_candidates[0], tgt_candidates[0]
                relations.append({
                    "type": "TRANSFERRED_TO",
                    "source": source,
                    "target": target,
                    "sourceText": source_text,
                    "confidence": 1.0,
                    "sourceRow": row_number,
                    "sourceFields": [f for f in normalized if "account" in f.casefold() or "sender" in f.casefold() or "receiver" in f.casefold()]
                })
                events.append({
                    "type": "TRANSACTION",
                    "date": timestamp,
                    "location": tower_loc,
                    "participants": [source, target],
                    "sourceText": source_text,
                    "confidence": 1.0,
                    "sourceRow": row_number
                })

        # 3. Vehicle Movement & Locations
        vehicle_mentions = typed.get("VEHICLE", [])
        loc_mentions = typed.get("LOCATION", [])
        person_mentions = typed.get("PERSON", [])
        
        if vehicle_mentions and loc_mentions:
            veh = vehicle_mentions[0]
            loc = loc_mentions[0]
            relations.append({
                "type": "VISITED",
                "source": veh,
                "target": loc,
                "sourceText": source_text,
                "confidence": 1.0,
                "sourceRow": row_number,
                "sourceFields": [f for f in normalized if _field_type(f.rsplit(".", 1)[-1]) in {"VEHICLE", "LOCATION"}]
            })
            participants = [veh]
            if person_mentions:
                drv = person_mentions[0]
                participants.append(drv)
                relations.append({
                    "type": "OPERATED",
                    "source": drv,
                    "target": veh,
                    "sourceText": source_text,
                    "confidence": 1.0,
                    "sourceRow": row_number,
                    "sourceFields": [f for f in normalized if _field_type(f.rsplit(".", 1)[-1]) in {"PERSON", "VEHICLE"}]
                })
            events.append({
                "type": "MOVEMENT",
                "date": timestamp,
                "location": loc,
                "participants": participants,
                "sourceText": source_text,
                "confidence": 1.0,
                "sourceRow": row_number
            })

        # 4. Surveillance / Activity with Persons and Locations
        elif person_mentions and loc_mentions and not caller and not sender:
            person = person_mentions[0]
            loc = loc_mentions[0]
            relations.append({
                "type": "VISITED",
                "source": person,
                "target": loc,
                "sourceText": source_text,
                "confidence": 0.95,
                "sourceRow": row_number,
                "sourceFields": [f for f in normalized if _field_type(f.rsplit(".", 1)[-1]) in {"PERSON", "LOCATION"}]
            })
            events.append({
                "type": "OBSERVATION",
                "date": timestamp,
                "location": loc,
                "participants": person_mentions,
                "sourceText": source_text,
                "confidence": 0.95,
                "sourceRow": row_number
            })

    structured_summary = {
        "records": normalized_records,
        "row_count": len(records),
        "column_count": len(columns),
        "columns": columns,
        "sheets": sheets_data or {}
    }
    return entities, relations, events, structured_summary


def process_structured(path: Path, mime_type: str) -> dict[str, Any]:
    suffix = path.suffix.lower()
    sheets_data = None
    if suffix == ".csv":
        records, warnings, columns = _records_from_csv(path)
    elif suffix in {".xls", ".xlsx"}:
        records, warnings, columns, sheets_data = _records_from_workbook(path)
    elif suffix == ".json":
        records, warnings, columns = _records_from_json(path)
    elif suffix == ".xml":
        records, warnings, columns = _records_from_xml(path)
    else:
        raise ValueError("Unsupported structured format")
        
    entities, relations, events, structured = _mentions(records, columns, sheets_data)
    structured["file_type"] = suffix.lstrip(".").upper()
    return {
        "text": "\n".join(json.dumps(record, ensure_ascii=False) for record in structured["records"]),
        "pages": [],
        "language": {"code": "unknown", "confidence": 0.0},
        "entities": entities,
        "categories": {},
        "relations": relations,
        "events": events,
        "warnings": warnings,
        "ingestion_method": "structured",
        "structured": structured
    }

